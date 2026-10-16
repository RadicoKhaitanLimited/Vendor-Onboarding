from django.core.management.base import BaseCommand, CommandError
from openpyxl import load_workbook

from apps.onboarding.models import (
    CompanyCodeMaster,
    CustomerCompanyCodeMaster,
    DistributionChannelMaster,
    DivisionMaster,
    SalesOrganizationMaster,
    TransportationZoneMaster,
)


class Command(BaseCommand):
    help = (
        'One-time load of customer-side sales-area master data: Division '
        '(from a Division List export), Transportation Zone (from a '
        'Transportation Zone export), Sales Organization + Distribution '
        'Channel (from the "Customer Sales Area" sheet of a ZBP_CUST-style '
        'export), and Customer Company Code (copied from the existing '
        'CompanyCodeMaster, since SAP company codes are shared legal '
        'entities, not vendor/customer-specific). Existing rows are left '
        'untouched; only missing codes are inserted.'
    )

    def add_arguments(self, parser):
        parser.add_argument('--division-xlsx', required=True)
        parser.add_argument('--transportation-zone-xlsx', required=True)
        parser.add_argument('--customer-sales-area-xlsx', required=True)

    def handle(self, *args, **options):
        self._load_divisions(options['division_xlsx'])
        self._load_transportation_zones(options['transportation_zone_xlsx'])
        self._load_sales_org_and_distribution_channel(options['customer_sales_area_xlsx'])
        self._copy_company_codes()

    def _open(self, path, sheet=None):
        try:
            wb = load_workbook(path, data_only=True, read_only=True)
        except FileNotFoundError:
            raise CommandError(f'File not found: {path}')
        return wb[sheet] if sheet else wb.worksheets[0]

    def _load_divisions(self, path):
        ws = self._open(path, 'Divisions')
        created = 0
        for code, name in list(ws.iter_rows(values_only=True))[1:]:
            if not code:
                continue
            _, was_created = DivisionMaster.objects.get_or_create(
                division=str(code).strip(),
                defaults={'description': str(name).strip() if name else ''},
            )
            created += was_created
        self.stdout.write(self.style.SUCCESS(f'Divisions: {created} created.'))

    def _load_transportation_zones(self, path):
        ws = self._open(path, 'Sheet1')
        created = 0
        for code, desc in list(ws.iter_rows(values_only=True))[1:]:
            if not code:
                continue
            _, was_created = TransportationZoneMaster.objects.get_or_create(
                zone=str(code).strip(),
                defaults={'description': str(desc).strip() if desc else ''},
            )
            created += was_created
        self.stdout.write(self.style.SUCCESS(f'Transportation zones: {created} created.'))

    def _load_sales_org_and_distribution_channel(self, path):
        ws = self._open(path, 'Customer Sales Area')
        rows = list(ws.iter_rows(values_only=True))[2:]  # skip title + header rows
        sorg_created = dchl_created = 0
        for row in rows:
            if not row or row[0] is None:
                continue
            s_code, s_name, d_code, d_name = row[0], row[1], row[2], row[3]
            s_code = str(s_code).strip()
            s_name = str(s_name).strip() if s_name else ''
            _, was_created = SalesOrganizationMaster.objects.get_or_create(
                sales_organization=s_code, defaults={'description': s_name},
            )
            sorg_created += was_created

            if d_code is not None:
                d_code = str(d_code).strip()
                d_name = str(d_name).strip() if d_name else ''
                _, was_created = DistributionChannelMaster.objects.get_or_create(
                    distribution_channel=d_code, defaults={'description': d_name},
                )
                dchl_created += was_created
        self.stdout.write(self.style.SUCCESS(
            f'Sales organizations: {sorg_created} created. '
            f'Distribution channels: {dchl_created} created.'
        ))

    def _copy_company_codes(self):
        created = 0
        for cc in CompanyCodeMaster.objects.all():
            _, was_created = CustomerCompanyCodeMaster.objects.get_or_create(
                company_code=cc.company_code, defaults={'name': cc.name},
            )
            created += was_created
        self.stdout.write(self.style.SUCCESS(f'Customer company codes: {created} created.'))
